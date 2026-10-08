"""Stream processed private CSVs to Excel without holding the minute panel in memory."""
import argparse
import csv
import json
import warnings
from datetime import date, datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo

SOURCE=Path(__file__).resolve().parent.parent
HEADERS=json.loads((SOURCE/'schemas/data_headers.json').read_text())
DICTIONARY=json.loads((SOURCE/'schemas/field_dictionary.json').read_text())


def parse_value(text, field):
    if text == '':return None
    if field in ('subject_id','activity_id','source_file','activity_type','activity_name','spo2_measurement_mode','status','error_type'):
        return text
    if text in ('True', 'False'):return text == 'True'
    if field in ('local_date','period_start','period_end','period_end_exclusive'):
        return date.fromisoformat(text)
    if field.endswith('_local') or field.endswith('_utc'):
        value=datetime.fromisoformat(text)
        if value.tzinfo:
            value=value.astimezone(timezone.utc if field.endswith('_utc') else ZoneInfo('Asia/Shanghai')).replace(tzinfo=None)
        return value
    try:
        value=float(text)
        return int(value) if value.is_integer() else value
    except ValueError:return text


def make_cell(sheet, value, header=False):
    cell=WriteOnlyCell(sheet,value=value)
    if isinstance(value,str):cell.data_type='s'
    if header:
        cell.font=Font(bold=True,color='FFFFFF',size=10)
        cell.fill=PatternFill('solid',fgColor='286482')
        cell.alignment=Alignment(wrap_text=True,vertical='center')
    elif isinstance(value,datetime):cell.number_format='yyyy-mm-dd hh:mm:ss'
    elif isinstance(value,date):cell.number_format='yyyy-mm-dd'
    return cell


def add_csv(workbook,path,name,expected_headers=None):
    with path.open(newline='',encoding='utf-8-sig') as stream:
        reader=csv.DictReader(stream)
        headers=reader.fieldnames
        if not headers:return
        if expected_headers is not None and headers != expected_headers:
            raise ValueError('Unexpected CSV schema: '+name)
        sheet=workbook.create_sheet(name);sheet.freeze_panes='B3';sheet.sheet_view.showGridLines=False
        sheet.merged_cells.add(f'A1:{get_column_letter(min(6,len(headers)))}1')
        sheet.row_dimensions[1].height=32;sheet.row_dimensions[2].height=64
        for index,field in enumerate(headers,1):sheet.column_dimensions[get_column_letter(index)].width=26
        sheet.append([make_cell(sheet,name+' | 每条记录保留受访者编号和时间，字段说明见README',True)])
        sheet.append([make_cell(sheet,h,True) for h in headers])
        count=0
        for record in reader:
            sheet.append([make_cell(sheet,parse_value(record[h],h)) for h in headers]);count+=1
        if count:
            ref=f'A2:{get_column_letter(len(headers))}{count+2}'
            table=Table(displayName=name+'_table',ref=ref,
                        tableColumns=[TableColumn(id=i,name=h) for i,h in enumerate(headers,1)],
                        autoFilter=AutoFilter(ref=ref))
            table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
            with warnings.catch_warnings():
                warnings.filterwarnings('ignore',message='In write-only mode you must add table columns manually')
                sheet.add_table(table)
        print(name,count)


def readme(workbook,config,quality=False):
    sheet=workbook.create_sheet('README');sheet.freeze_panes='A2'
    for col,width in zip('ABCDE',[24,38,34,24,66]):sheet.column_dimensions[col].width=width
    rows=[['Garmin数据缺失检查说明' if quality else 'Garmin手环数据字段与来源说明'],
          [f"观测期：{config['start']} 至 {config['end_exclusive']}，右端不含，Asia/Shanghai。"],
          ['第1行是工作表说明，第2行是变量名，第3行起是数据；pandas读取时skiprows=1。'],
          ['官网心率按固定采样槽位计算覆盖率；接口失败不当作100%缺失。具体阈值与采样间隔见run_config.json。'],
          ['FIT分钟记录只保留真实观测，不插值、不填生理值0；同分钟有效读数取均值，计数为0表示无有效观测。'],
          ['官方日值用于总步数、距离、活动热量、强度分钟、楼梯高度和身体电量；日级心率统计保留FIT口径。'],
          ['睡眠按醒来日合并，sleep_record_valid为项目质控；未公开字段不是通用官方定义，confidence不是准确率。'],
          ['数据源：原始FIT → SDK解码CSV → 清洗聚合；日值/睡眠/活动摘要来自登录后的Connect非公开JSON接口。'],
          ['官网槽位缺失率与FIT分钟空白率不是同一口径，均不能直接作为未佩戴时间比例。']]
    for i,row in enumerate(rows,1):
        sheet.row_dimensions[i].height=48;sheet.merged_cells.add(f'A{i}:E{i}')
        cells=[make_cell(sheet,v,i==1) for v in row]
        for cell in cells:cell.alignment=Alignment(wrap_text=True,vertical='center')
        sheet.append(cells)
    sheet.append([make_cell(sheet,v,True) for v in ['工作表','变量名','中文含义','单位/格式','数据来源与处理']])
    if quality:
        quality_fields=[
          ['subject_missingness','status','纳入、排除、登录失败或初筛失败','文本','所有授权且佩戴手环的对象均保留一行'],
          ['subject_missingness','heart_rate_missing_rate','官网心率初筛缺失率','比例','1−有效采样槽位/预期槽位'],
          ['subject_missingness','estimated_covered_minutes','按官网槽位估计覆盖时间','分钟','有效槽位×采样秒数/60，不新增观测'],
          ['heart_rate_screen_by_day','request_success','该日心率接口是否成功','布尔值','失败表示缺失率未知'],
          ['data_missingness_by_day','fit_hr_minute_missing_rate','FIT心率分钟空白率','比例','1−有有效心率的分钟/1440'],
          ['data_missingness_by_day','stress_minute_missing_rate','压力分钟空白率','比例','1−有有效压力的分钟/1440'],
          ['data_missingness_by_day','spo2_minute_missing_rate','血氧分钟空白率','比例','1−有有效血氧的分钟/1440'],
          ['download_status','status','接口下载状态','文本','downloaded/empty/error；empty与error含义不同'],
          ['download_status','detail','错误类型','文本','只记录类型，不输出密码或登录凭据'],
        ]
        for row in quality_fields:sheet.append(row)
    else:
        for item in DICTIONARY:
            row=[item[k] for k in ['sheet','field','meaning_zh','unit','source_and_rule_zh']]
            cells=[make_cell(sheet,v) for v in row]
            for cell in cells:cell.alignment=Alignment(wrap_text=True,vertical='center')
            sheet.append(cells)


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--root',required=True,type=Path);args=parser.parse_args()
    root=args.root.resolve();config=json.loads((root/'run_config.json').read_text())
    output=root/'results';workbook=Workbook(write_only=True);workbook.properties.creator='Garmin data pipeline'
    readme(workbook,config)
    for name,headers in HEADERS.items():add_csv(workbook,output/f'{name}.csv',name,headers)
    workbook.save(output/'dataset.xlsx')
    quality=Workbook(write_only=True);quality.properties.creator='Garmin data pipeline';readme(quality,config,True)
    for name in ['subject_missingness','heart_rate_screen_by_day','data_missingness_by_day','processing_issues']:
        path=output/f'{name}.csv'
        if path.exists():add_csv(quality,path,name)
    path=root/'logs/download_status.csv'
    if path.exists():add_csv(quality,path,'download_status')
    quality.save(output/'missingness.xlsx')


if __name__=='__main__':main()
