"""Offline tests use invented readings; no real accounts, dates or downloaded records."""
import ast
import csv
import json
import os
import subprocess
import sys
import tempfile
import unittest
import zipfile
from datetime import date, datetime, timezone
from pathlib import Path
from xml.etree import ElementTree as ET

PUBLIC = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PUBLIC / 'src'))
import pandas as pd
from openpyxl import load_workbook
import download
import process
import export_excel
import privacy_check
from daily_summary import attribution_date
from fit2csv import resolve_timestamp_16
from garmin_fit_sdk import convert_timestamp_to_datetime, convert_datetime_to_timestamp


def write_json(path, obj):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj))


def write_csv(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(rows).to_csv(path, index=False)


class ScreeningTests(unittest.TestCase):
    def test_fixed_slots_missing_omitted_and_duplicate(self):
        start = int(datetime.fromisoformat('2030-01-01T00:00:00+08:00').timestamp()*1000)
        row = download.screen_day({'heartRateValues': [
            [start, 60], [start, 60], [start+30000, 70], [start+120000, None],
            [start+240000, 255], [start+360000, 0], [start-1, 80], [start+86400000, 80]
        ]}, '2030-01-01', 120)
        self.assertEqual(row['expected_slots'], 720)
        self.assertEqual(row['valid_slots'], 1)
        self.assertAlmostEqual(row['heart_rate_missing_rate'], 719/720)

    def test_threshold_is_inclusive(self):
        self.assertTrue(download.excluded_for_missing(144, 720, .8))
        self.assertFalse(download.excluded_for_missing(145, 720, .8))
        self.assertTrue(download.excluded_for_missing(143, 720, .8))

    def test_subject_path_traversal(self):
        for sid in ['../data', '/root', 'a/b', '', 'a b']:
            with self.assertRaises(ValueError):download.validate_subject_id(sid)
        self.assertEqual(download.validate_subject_id('SAMPLE_001'), 'SAMPLE_001')

    def test_empty_response_is_missing(self):
        self.assertEqual(download.screen_day({}, '2030-01-01', 120)['heart_rate_missing_rate'], 1)

    def test_private_root_must_be_outside_public(self):
        with tempfile.TemporaryDirectory() as tmp:
            accounts=Path(tmp)/'accounts.json';accounts.write_text('[]')
            command=[sys.executable, str(PUBLIC/'src/download.py'), '--accounts', str(accounts),
                     '--root', str(PUBLIC), '--start', '2030-01-01', '--end', '2030-01-02']
            r=subprocess.run(command,capture_output=True,text=True)
            self.assertNotEqual(r.returncode, 0)
            self.assertIn('outside',r.stderr)


class AggregationTests(unittest.TestCase):
    def setUp(self):
        process.START=pd.Timestamp('2030-01-01',tz='Asia/Shanghai')
        process.STOP=process.START+pd.Timedelta(days=1)
        process.DAYS=['2030-01-01']

    def test_minute_panel_deduplicates_and_preserves_empty(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)
            write_csv(root/'csv/monitoring.csv',[
                {'local_time':'2030-01-01T00:00:10+08:00','heart_rate':60,'source_file':'a.fit'},
                {'local_time':'2030-01-01T00:00:10+08:00','heart_rate':60,'source_file':'b.fit'},
                {'local_time':'2030-01-01T00:00:40+08:00','heart_rate':80,'source_file':'a.fit'},
                {'local_time':'2030-01-01T00:01:00+08:00','heart_rate':0,'source_file':'a.fit'},
                {'local_time':'2030-01-02T00:00:00+08:00','heart_rate':100,'source_file':'a.fit'}])
            write_csv(root/'csv/stress_level.csv',[
                {'local_time':'2030-01-01T00:00:00+08:00','stress_level_value':0,'source_file':'a.fit'},
                {'local_time':'2030-01-01T00:00:30+08:00','stress_level_value':-1,'source_file':'a.fit'}])
            write_csv(root/'csv/spo2_data.csv',[
                {'local_time':'2030-01-01T00:00:10+08:00','reading_spo2':96,'source_file':'a.fit'},
                {'local_time':'2030-01-01T00:00:10+08:00','reading_spo2':96,'source_file':'b.fit'},
                {'local_time':'2030-01-01T00:00:50+08:00','reading_spo2':98,'source_file':'a.fit'}])
            panel=process.panel(root,'SAMPLE_001')
            self.assertEqual(len(panel),1440)
            first=panel.iloc[0];second=panel.iloc[1]
            self.assertEqual(first.heart_rate_mean_bpm,70)
            self.assertEqual(first.heart_rate_observation_count,2)
            self.assertEqual(first.stress_score,0)
            self.assertEqual(first.stress_valid_observation_count,1)
            self.assertEqual(first.stress_raw_observation_count,2)
            self.assertEqual(first.spo2_percent,97)
            self.assertTrue(pd.isna(second.heart_rate_mean_bpm))
            self.assertEqual(second.heart_rate_observation_count,0)
            self.assertTrue(pd.isna(second.spo2_percent))

    def test_midnight_shift_only_cumulative(self):
        result=attribution_date(pd.Series(['2030-01-02T00:00:00+08:00','2030-01-02T00:00:01+08:00']))
        self.assertEqual(result.tolist(),['2030-01-01','2030-01-02'])

    def test_compressed_timestamp_wrap(self):
        base=65534
        rows=[{'timestamp':convert_timestamp_to_datetime(base)},{'timestamp_16':1},{'timestamp_16':4}]
        self.assertEqual(resolve_timestamp_16(rows),2)
        self.assertEqual(convert_datetime_to_timestamp(rows[1]['timestamp']),65537)
        self.assertEqual(convert_datetime_to_timestamp(rows[2]['timestamp']),65540)

    def test_daily_summary_retains_invalid_stress_denominator(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)
            write_csv(root/'monitoring.csv',[{'subject_id':'SAMPLE_001','source_file':'a.fit',
                'local_time':'2030-01-01T12:00:00+08:00','local_date':'2030-01-01',
                'heart_rate':60,'activity_type':'walking','cycles':4,'distance':5,'active_calories':0}])
            write_csv(root/'stress_level.csv',[{'subject_id':'SAMPLE_001','source_file':'a.fit',
                'local_time':'2030-01-01T12:00:00+08:00','local_date':'2030-01-01','stress_level_value':-1}])
            r=subprocess.run([sys.executable,str(PUBLIC/'src/daily_summary.py'),str(root),'-o',str(root/'daily.csv')],capture_output=True,text=True)
            self.assertEqual(r.returncode,0,r.stderr)
            row=pd.read_csv(root/'daily.csv').iloc[0]
            self.assertEqual(row.stress_total,1)
            self.assertEqual(row.stress_n,0)
            self.assertEqual(row.stress_valid_pct,0)
            self.assertTrue(pd.isna(row.stress_mean))


class ExportTests(unittest.TestCase):
    def test_value_types_and_timezones(self):
        self.assertIsNone(export_excel.parse_value('','heart_rate_mean_bpm'))
        self.assertEqual(export_excel.parse_value('0','stress_score'),0)
        self.assertEqual(export_excel.parse_value('000123','activity_id'),'000123')
        self.assertEqual(export_excel.parse_value('True','sleep_record_valid'),True)
        self.assertEqual(export_excel.parse_value('2030-01-01T00:00:00+08:00','spo2_reading_time_utc'),datetime(2029,12,31,16))
        self.assertEqual(export_excel.parse_value('2029-12-31T16:00:00Z','minute_start_local'),datetime(2030,1,1))

    def test_small_offline_end_to_end(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);sid='SAMPLE_001';sub=root/'data'/sid
            write_json(root/'run_config.json',{'start':'2030-01-01','end_exclusive':'2030-01-02'})
            write_json(root/'screening'/f'{sid}_summary.json',{'subject_id':sid,'status':'included','download_phase':'finished'})
            write_json(root/'screening'/f'{sid}_daily.json',[{'subject_id':sid,'local_date':'2030-01-01','request_success':True,'expected_slots':720,'valid_slots':720}])
            write_csv(sub/'csv/monitoring.csv',[{'local_time':'2030-01-01T12:00:00+08:00','heart_rate':60,'source_file':'a.fit'}])
            write_json(sub/'official/json/2030-01-01.json',{'totalSteps':0,'activeSeconds':10,'highlyActiveSeconds':20})
            write_json(sub/'raw/activities/activities.json',[{'activityId':'000123','startTimeLocal':'2030-01-01 12:00:00',
                  'activityName':'=1+1','distance':1000,'duration':300,'averageSpeed':2}])
            for tool in ['process','export_excel']:
                r=subprocess.run([sys.executable,str(PUBLIC/'src'/f'{tool}.py'),'--root',str(root)],capture_output=True,text=True)
                self.assertEqual(r.returncode,0,r.stdout+r.stderr)
            book=load_workbook(root/'results/dataset.xlsx',read_only=False)
            self.assertEqual(book.sheetnames,['README','daily_data','minute_data','activities'])
            self.assertEqual(book['daily_data'].max_row,3)
            self.assertEqual(book['minute_data'].max_row,1442)
            self.assertEqual(book['activities']['B3'].value,'000123')
            self.assertEqual(book['activities']['F3'].value,'=1+1')
            self.assertEqual(book['activities']['F3'].data_type,'s')
            self.assertIsInstance(book['minute_data']['C3'].value,datetime)
            for name in ['daily_data','minute_data','activities']:
                self.assertEqual(book[name].freeze_panes,'B3')
                self.assertTrue(next(iter(book[name].tables.values())).autoFilter.ref)
            headers=[c.value for c in book['daily_data'][2]]
            self.assertEqual(book['daily_data'].cell(3,headers.index('official_total_steps')+1).value,0)
            book.close()
            book=load_workbook(root/'results/missingness.xlsx',read_only=True)
            self.assertEqual(sum(1 for _ in book['subject_missingness'].iter_rows()),3)
            book.close()


class PrivacyTests(unittest.TestCase):
    def test_release_allowlist(self):
        self.assertTrue(privacy_check.allowed('examples/minute_data_headers.csv'))
        for name in ['results/a.csv','src/data.csv','docs/screenshot.png','accounts.json','src/.tokens/cache.json']:
            self.assertFalse(privacy_check.allowed(name))

    def test_sensitive_patterns(self):
        for text in ['/'+'Users'+'/someone/data','13'+'123456789','AB'+'CD'+'1234','somebody'+'@'+'example.com']:
            self.assertTrue(privacy_check.inspect_text(text))
        self.assertEqual(privacy_check.inspect_text('SAMPLE_001 GARMIN_ACCOUNT_001'),[])

    def test_multiline_example_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'examples/minute_data_headers.csv';path.parent.mkdir()
            path.write_text('subject_id,value\nSAMPLE_001,60\n')
            r=subprocess.run([sys.executable,str(PUBLIC/'src/privacy_check.py'),'--root',tmp],capture_output=True,text=True)
            self.assertEqual(r.returncode,1)
            self.assertIn('example_csv_contains_records',r.stdout)

    def test_scanner_checks_staged_content(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)
            subprocess.run(['git','init','-q',tmp],check=True)
            p=root/'README.md';p.write_text('13'+'123456789')
            subprocess.run(['git','-C',tmp,'add','README.md'],check=True)
            p.write_text('safe working copy')
            r=subprocess.run([sys.executable,str(PUBLIC/'src/privacy_check.py'),'--root',tmp,'--staged'],capture_output=True,text=True)
            self.assertEqual(r.returncode,1)
            self.assertIn('mobile_number',r.stdout)

    def test_public_python_parses(self):
        for path in (PUBLIC/'src').glob('*.py'):ast.parse(path.read_text())


if __name__ == '__main__':unittest.main()
